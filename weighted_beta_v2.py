import yfinance as yf
import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime
import warnings

input_file = "Holdings.xlsx"
warnings.filterwarnings('ignore')

# ---------------------------------------------------------------------------
# CONFIG
# ---------------------------------------------------------------------------

TICKER_REMAP = {
    ("PRX", "EUR"): ("PROSY", "USD"),
}

FX_SHEET   = "FX Rates"
FX_ROW_USD = 3
FX_ROW_CNH = 4

FX_RATES = {
    "USD": 7.78,
    "CNH": 1.071,
    "HKD": 1.0,
}

BLUE  = "FF0000FF"
BLACK = "FF000000"
GREEN = "FF008000"

HEADER_FILL = PatternFill("solid", start_color="FF1F4E79")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFFFF", size=9)
DATA_FONT   = Font(name="Arial", size=9)
INPUT_FONT  = Font(name="Arial", size=9, color=BLUE)
FORMULA_FONT= Font(name="Arial", size=9, color=BLACK)
LINK_FONT   = Font(name="Arial", size=9, color=GREEN)

THIN   = Side(style="thin", color="FFD9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# ---------------------------------------------------------------------------
# TICKER FORMATTERS FOR YFINANCE
# ---------------------------------------------------------------------------

def yf_ticker(ticker: str, currency: str) -> str:
    """Convert a holding ticker + currency to a yfinance symbol."""
    t = str(ticker).strip()
    # Excel reads numeric tickers as floats (e.g. 1008 → "1008.0"); strip the .0
    if t.endswith(".0") and t[:-2].isdigit():
        t = t[:-2]
    if currency == "CNH" and len(t) == 6:
        # A-share: Shanghai starts with 6, Shenzhen with 0/3, Beijing with 4/8
        if t.startswith("6"):
            return f"{t}.SS"
        elif t.startswith("4") or t.startswith("8"):
            return f"{t}.BJ"
        else:
            return f"{t}.SZ"
    elif currency in ("HKD", "CNH"):
        # HK: zero-pad to 4 digits + .HK
        return f"{t.zfill(4)}.HK"
    else:
        # USD / EUR — use symbol as-is (TICKER_REMAP already applied)
        return t

# ---------------------------------------------------------------------------
# PRICE FETCHER
# ---------------------------------------------------------------------------

def get_prices(ticker: str, currency: str) -> pd.DataFrame | None:
    ticker, currency = TICKER_REMAP.get((ticker, currency), (ticker, currency))
    sym = yf_ticker(ticker, currency)
    for period in ("3mo", "6mo", "1y"):
        try:
            # Ticker.history() is more reliable than download() for single tickers
            hist = yf.Ticker(sym).history(period=period, auto_adjust=True)
            if hist is None or hist.empty:
                continue
            close = hist["Close"].dropna()
            if close.empty:
                continue
            if period != "3mo":
                print(f"  [{sym}] no data for 3mo — fetched with period={period}")
            result = pd.DataFrame({"Close": close})
            result.index = pd.to_datetime(result.index).tz_localize(None)
            result.index.name = "Date"
            return result
        except Exception as e:
            print(f"  yfinance prices failed for {sym} (period={period}): {e}")
    print(f"  yfinance prices: no data for {sym} across all periods tried")
    return None

# ---------------------------------------------------------------------------
# FUNDAMENTALS FETCHER
# ---------------------------------------------------------------------------

def _safe_float(val, allow_zero: bool = False) -> float | None:
    try:
        f = float(val)
        if pd.isna(f):
            return None
        if f == 0 and not allow_zero:
            return None
        return f
    except Exception:
        return None


def get_fundamentals(ticker: str, currency: str) -> dict:
    """
    Fetch PE, PB, dividend yield, and market cap via yfinance.
    Market cap is converted to HKD using FX_RATES.
    """
    ticker, currency = TICKER_REMAP.get((ticker, currency), (ticker, currency))
    sym = yf_ticker(ticker, currency)

    result = {
        "pe": None, "pb": None, "div_yield": None,
        "mkt_cap_local": None, "mkt_cap_hkd": None,
        "gics_sector": None, "gics_industry": None,
        "source": f"yfinance ({sym})",
    }

    FIELDS = {
        "pe":            "trailingPE",
        "pb":            "priceToBook",
        "div_yield":     "trailingAnnualDividendYield",
        "mkt_cap_local": "marketCap",
    }

    try:
        info = yf.Ticker(sym).info

        # Detect empty/delisted response — yfinance returns a minimal stub with only
        # a few keys (e.g. just "symbol", "trailingPegRatio") for unknown/delisted tickers
        if len(info) < 10:
            print(f"  [{sym}] yfinance returned a near-empty response "
                  f"({len(info)} keys: {list(info.keys())}) — "
                  f"ticker may be delisted, suspended, or unrecognised.")
            return result

        any_na = False
        for key, yf_key in FIELDS.items():
            raw = info.get(yf_key)
            # ETFs have totalAssets instead of marketCap
            if key == "mkt_cap_local" and raw is None:
                raw = info.get("totalAssets")
            parsed = _safe_float(raw, allow_zero=(key == "div_yield"))
            result[key] = parsed
            if parsed is None:
                reason = "key absent" if yf_key not in info else f"raw={raw!r}"
                print(f"  [{sym}] {key}: N/A ({reason})")
                any_na = True

        if any_na:
            relevant = {k: info[k] for k in
                        ["quoteType", "exchange", "marketCap", "totalAssets",
                         "trailingPE", "forwardPE", "priceToBook", "dividendYield",
                         "enterpriseValue", "bookValue"]
                        if k in info}
            print(f"  [{sym}] available valuation fields: {relevant}")

        result["gics_sector"]   = info.get("sector")   or None
        result["gics_industry"] = info.get("industry") or None

    except Exception as e:
        print(f"  yfinance fundamentals failed for {sym}: {e}")
        return result

    fx = FX_RATES.get(currency, 1.0)
    if result["mkt_cap_local"] is not None:
        result["mkt_cap_hkd"] = round(result["mkt_cap_local"] * fx)

    return result

# ---------------------------------------------------------------------------
# TRADING-DAY ALIGNMENT
# ---------------------------------------------------------------------------

def align_to_benchmark(asset: pd.DataFrame, benchmark: pd.DataFrame) -> pd.DataFrame:
    asset_aligned = asset.reindex(benchmark.index).ffill()
    combined = pd.DataFrame({
        "asset_close": asset_aligned["Close"],
        "bm_close":    benchmark["Close"],
    }).dropna()
    return combined

# ---------------------------------------------------------------------------
# EXCEL HELPERS
# ---------------------------------------------------------------------------

def col(n: int) -> str:
    return get_column_letter(n)

def write_header_row(ws, headers: list[str], row: int = 1):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.font      = HEADER_FONT
        cell.fill      = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
        cell.border    = BORDER
        ws.column_dimensions[col(c)].width = max(14, len(h) + 2)

def style_data_cell(cell, font=None, num_fmt=None):
    cell.font      = font or DATA_FONT
    cell.border    = BORDER
    cell.alignment = Alignment(horizontal="right")
    if num_fmt:
        cell.number_format = num_fmt

# ---------------------------------------------------------------------------
# FX RATES SHEET
# ---------------------------------------------------------------------------

def create_fx_sheet(wb: Workbook):
    ws = wb.create_sheet(FX_SHEET, 0)
    ws.sheet_view.showGridLines = False

    ws["A1"] = "FX Rate Assumptions (vs HKD)"
    ws["A1"].font = Font(name="Arial", bold=True, size=11)
    ws.merge_cells("A1:D1")

    headers = ["Currency Pair", "Rate (to HKD)", "Last Updated"]
    write_header_row(ws, headers, row=2)

    today_str = datetime.today().strftime("%d-%b-%Y")
    rows = [
        ("USD/HKD", 7.78,  today_str),
        ("CNH/HKD", 1.071, today_str),
    ]
    for r, (pair, rate, date) in enumerate(rows, 3):
        ws.cell(row=r, column=1, value=pair).font = DATA_FONT
        rate_cell = ws.cell(row=r, column=2, value=rate)
        rate_cell.font          = INPUT_FONT
        rate_cell.number_format = "0.0000"
        ws.cell(row=r, column=3, value=date).font = DATA_FONT
        for c in range(1, 4):
            ws.cell(row=r, column=c).border = BORDER

    yellow = PatternFill("solid", start_color="FFFFFF00")
    ws["B3"].fill = yellow
    ws["B4"].fill = yellow

    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 50
    return ws

# ---------------------------------------------------------------------------
# FX FORMULA HELPERS
# ---------------------------------------------------------------------------

def fx_formula(currency: str) -> str:
    if currency == "USD":
        return f"='{FX_SHEET}'!$B${FX_ROW_USD}"
    elif currency == "CNH":
        return f"='{FX_SHEET}'!$B${FX_ROW_CNH}"
    else:
        return "=1"

def fx_formula_for_mktcap(currency: str) -> str:
    if currency in ("USD", "EUR"):
        return f"='{FX_SHEET}'!$B${FX_ROW_USD}"
    elif currency == "CNH":
        return f"='{FX_SHEET}'!$B${FX_ROW_CNH}"
    else:
        return "=1"

# ---------------------------------------------------------------------------
# WRITE STOCK SHEET
# ---------------------------------------------------------------------------

def write_stock_sheet(wb: Workbook, ticker: str, orig_currency: str,
                      combined: pd.DataFrame, beta_result: dict):
    ticker, resolved_ccy = TICKER_REMAP.get((ticker, orig_currency), (ticker, orig_currency))

    sheet_name = ticker[:31].replace("/", "_").replace("\\", "_")
    ws = wb.create_sheet(sheet_name)
    ws.sheet_view.showGridLines = False

    headers = [
        "Date",
        f"{ticker} Close ({resolved_ccy})",
        "FX Rate (to HKD)",
        f"{ticker} Close (HKD)",
        "2800.HK Close (HKD)",
        f"{ticker} Return (HKD)",
        "2800.HK Return",
    ]
    write_header_row(ws, headers, row=1)
    ws.column_dimensions["A"].width = 13

    data_rows = combined.reset_index()

    for r_idx, data_row in enumerate(data_rows.itertuples(index=False), start=2):
        date_val  = data_row[0]
        asset_val = data_row[1]
        bm_val    = data_row[2]

        c = ws.cell(row=r_idx, column=1, value=date_val)
        c.number_format = "DD-MMM-YY"; style_data_cell(c, DATA_FONT)

        c = ws.cell(row=r_idx, column=2, value=round(float(asset_val), 4))
        style_data_cell(c, DATA_FONT, "0.0000")

        c = ws.cell(row=r_idx, column=3, value=fx_formula(resolved_ccy))
        style_data_cell(c, LINK_FONT, "0.0000")

        c = ws.cell(row=r_idx, column=4, value=f"={col(2)}{r_idx}*{col(3)}{r_idx}")
        style_data_cell(c, FORMULA_FONT, "0.0000")

        c = ws.cell(row=r_idx, column=5, value=round(float(bm_val), 4))
        style_data_cell(c, DATA_FONT, "0.0000")

        if r_idx == 2:
            ws.cell(row=r_idx, column=6, value="—"); style_data_cell(ws.cell(row=r_idx, column=6), DATA_FONT)
        else:
            formula = f"=({col(4)}{r_idx}-{col(4)}{r_idx-1})/{col(4)}{r_idx-1}"
            c = ws.cell(row=r_idx, column=6, value=formula)
            style_data_cell(c, FORMULA_FONT, "0.00%")

        if r_idx == 2:
            ws.cell(row=r_idx, column=7, value="—"); style_data_cell(ws.cell(row=r_idx, column=7), DATA_FONT)
        else:
            formula = f"=({col(5)}{r_idx}-{col(5)}{r_idx-1})/{col(5)}{r_idx-1}"
            c = ws.cell(row=r_idx, column=7, value=formula)
            style_data_cell(c, FORMULA_FONT, "0.00%")

    last_data_row    = len(data_rows) + 1
    summary_start    = last_data_row + 3
    last_return_row  = last_data_row
    first_return_row = last_return_row - 44

    labels = [
        ("45-Day Covariance (HKD returns):", beta_result["covariance"]),
        ("45-Day Benchmark Variance:",        beta_result["variance"]),
        ("FINAL 45-DAY BETA:",                beta_result["beta"]),
    ]

    if isinstance(beta_result["beta"], float):
        f_ret = first_return_row
        l_ret = last_return_row
        cov_formula  = (f"=_xlfn.COVARIANCE.S({col(6)}{f_ret}:{col(6)}{l_ret},"
                        f"{col(7)}{f_ret}:{col(7)}{l_ret})")
        var_formula  = f"=_xlfn.VAR.S({col(7)}{f_ret}:{col(7)}{l_ret})"
        beta_formula = f"={col(2)}{summary_start}/{col(2)}{summary_start+1}"
        excel_formulas = [cov_formula, var_formula, beta_formula]
    else:
        excel_formulas = [beta_result["covariance"],
                          beta_result["variance"],
                          beta_result["beta"]]

    for i, (label, _) in enumerate(labels):
        sr = summary_start + i
        lc = ws.cell(row=sr, column=1, value=label)
        lc.font = Font(name="Arial", bold=(i == 2), size=9)
        lc.border = BORDER

        vc = ws.cell(row=sr, column=2, value=excel_formulas[i])
        vc.font   = Font(name="Arial", bold=(i == 2), size=9, color=BLACK)
        vc.border = BORDER
        if i == 2:
            gold = PatternFill("solid", start_color="FFFFF2CC")
            lc.fill = gold; vc.fill = gold
        if isinstance(excel_formulas[i], float):
            vc.number_format = "0.0000"

    ws.freeze_panes = "A2"

# ---------------------------------------------------------------------------
# FUNDAMENTALS SHEET
# ---------------------------------------------------------------------------

def write_fundamentals_sheet(wb: Workbook, holdings: pd.DataFrame,
                              fundamentals_map: dict):
    ws = wb.create_sheet("Fundamentals")
    ws.sheet_view.showGridLines = False

    ws["A1"] = "Fundamental Metrics Snapshot"
    ws["A1"].font = Font(name="Arial", bold=True, size=13)
    ws.merge_cells("A1:L1")

    ws["A2"] = (f"Point-in-time snapshot | "
                f"Mkt Cap (HKD) recalculates with FX Rates sheet | "
                f"Generated: {datetime.today().strftime('%d-%b-%Y %H:%M')}")
    ws["A2"].font = Font(name="Arial", italic=True, size=9, color="FF555555")
    ws.merge_cells("A2:L2")

    headers = [
        "Ticker", "Currency", "GICS Sector", "GICS Industry",
        "P/E (Trailing)", "P/B",
        "Div Yield",
        "Mkt Cap (Local)", "FX Rate (→HKD)", "Mkt Cap (HKD)",
        "Fetched At",
    ]
    write_header_row(ws, headers, row=4)

    col_widths = [12, 10, 22, 28, 16, 10, 12, 20, 16, 22, 18]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[col(i)].width = w

    today_str = datetime.today().strftime("%d-%b-%Y %H:%M")

    for r_idx, (_, row) in enumerate(holdings.iterrows(), start=5):
        ticker   = str(row["Ticker"]).strip()
        currency = str(row["Local Currency"]).strip()
        ticker, currency = TICKER_REMAP.get((ticker, currency), (ticker, currency))

        f = fundamentals_map.get(ticker, {})
        pe           = f.get("pe")
        pb           = f.get("pb")
        div_yield    = f.get("div_yield")
        mkt_local    = f.get("mkt_cap_local")
        gics_sector  = f.get("gics_sector")
        gics_industry= f.get("gics_industry")

        def _val(v, fallback="N/A"):
            return v if v is not None else fallback

        # A: Ticker
        c = ws.cell(row=r_idx, column=1, value=ticker)
        style_data_cell(c, DATA_FONT); c.alignment = Alignment(horizontal="left")

        # B: Currency
        c = ws.cell(row=r_idx, column=2, value=currency)
        style_data_cell(c, DATA_FONT); c.alignment = Alignment(horizontal="center")

        # C: GICS Sector
        c = ws.cell(row=r_idx, column=3, value=_val(gics_sector))
        style_data_cell(c, DATA_FONT); c.alignment = Alignment(horizontal="left")

        # D: GICS Industry
        c = ws.cell(row=r_idx, column=4, value=_val(gics_industry))
        style_data_cell(c, DATA_FONT); c.alignment = Alignment(horizontal="left")

        # E: P/E
        c = ws.cell(row=r_idx, column=5, value=_val(pe))
        style_data_cell(c, DATA_FONT, "0.00" if pe is not None else None)

        # F: P/B
        c = ws.cell(row=r_idx, column=6, value=_val(pb))
        style_data_cell(c, DATA_FONT, "0.00" if pb is not None else None)

        # G: Div Yield
        c = ws.cell(row=r_idx, column=7, value=_val(div_yield))
        style_data_cell(c, DATA_FONT, '0.00%' if div_yield is not None else None)

        # H: Mkt Cap Local
        c = ws.cell(row=r_idx, column=8, value=_val(mkt_local))
        style_data_cell(c, DATA_FONT, "#,##0" if mkt_local is not None else None)

        # I: FX Rate
        fx_fml = fx_formula_for_mktcap(currency)
        c = ws.cell(row=r_idx, column=9, value=fx_fml)
        style_data_cell(c, LINK_FONT, "0.0000")

        # J: Mkt Cap HKD
        if mkt_local is not None:
            hkd_fml = f"={col(8)}{r_idx}*{col(9)}{r_idx}"
            c = ws.cell(row=r_idx, column=10, value=hkd_fml)
            style_data_cell(c, FORMULA_FONT, "#,##0")
        else:
            c = ws.cell(row=r_idx, column=10, value="N/A")
            style_data_cell(c, DATA_FONT)

        # K: Fetched At
        c = ws.cell(row=r_idx, column=11, value=today_str)
        style_data_cell(c, DATA_FONT); c.alignment = Alignment(horizontal="center")

    ws.freeze_panes = "A5"

# ---------------------------------------------------------------------------
# SUMMARY SHEET
# ---------------------------------------------------------------------------

def _weighted_avg(holdings: pd.DataFrame, beta_map: dict,
                  fundamentals_map: dict, key: str,
                  long_only: bool = False) -> float | None:
    """Compute weighted average of a fundamentals/beta key across holdings.
    long_only=True restricts to positive-weight (long) positions only.
    """
    total_w  = 0.0
    total_wv = 0.0
    for _, row in holdings.iterrows():
        ticker   = str(row["Ticker"]).strip()
        currency = str(row["Local Currency"]).strip()
        ticker, _ = TICKER_REMAP.get((ticker, currency), (ticker, currency))
        weight = pd.to_numeric(row["Weighting"], errors="coerce")
        if pd.isna(weight):
            continue
        if long_only and weight <= 0:
            continue
        if key == "beta":
            val = beta_map.get(ticker, {}).get("beta")
            if not isinstance(val, float):
                continue
        else:
            val = fundamentals_map.get(ticker, {}).get(key)
        if val is None or not isinstance(val, (int, float)):
            continue
        total_wv += weight * val
        total_w  += weight
    return round(total_wv / total_w, 6) if total_w > 0 else None


def _long_sector_exposure(holdings: pd.DataFrame,
                          fundamentals_map: dict) -> list[tuple[str, int, float]]:
    """
    Return [(sector, n_tickers, pct_of_long_book), ...] sorted descending by exposure.
    pct is expressed as a fraction (e.g. 0.30 = 30%) relative to total long weight.
    Positions with no sector data are grouped under 'Unknown / N/A'.
    """
    sector_weights: dict[str, float] = {}
    sector_counts:  dict[str, int]   = {}
    total_long = 0.0
    for _, row in holdings.iterrows():
        ticker   = str(row["Ticker"]).strip()
        currency = str(row["Local Currency"]).strip()
        ticker, _ = TICKER_REMAP.get((ticker, currency), (ticker, currency))
        weight = pd.to_numeric(row["Weighting"], errors="coerce")
        if pd.isna(weight) or weight <= 0:
            continue
        sector = fundamentals_map.get(ticker, {}).get("gics_sector") or "Unknown / N/A"
        sector_weights[sector] = sector_weights.get(sector, 0.0) + weight
        sector_counts[sector]  = sector_counts.get(sector, 0) + 1
        total_long += weight
    if total_long == 0:
        return []
    return sorted(
        [(s, sector_counts[s], w / total_long) for s, w in sector_weights.items()],
        key=lambda x: x[2], reverse=True,
    )


def write_summary_sheet(wb: Workbook, holdings: pd.DataFrame,
                        beta_map: dict, fundamentals_map: dict):
    ws = wb.create_sheet("Summary", 0)
    wb.active = ws
    ws.sheet_view.showGridLines = False

    ws["A1"] = "Beta Calculation Summary"
    ws["A1"].font = Font(name="Arial", bold=True, size=13)
    ws.merge_cells("A1:L1")

    ws["A2"] = (f"Benchmark: 2800.HK (Hang Seng ETF) | "
                f"Generated: {datetime.today().strftime('%d-%b-%Y %H:%M')}")
    ws["A2"].font = Font(name="Arial", italic=True, size=9, color="FF555555")
    ws.merge_cells("A2:L2")

    headers = [
        "Ticker", "Currency", "Weighting", "Long/Short", "Aligned Days",
        "45-Day Beta (HKD)", "Status", "GICS Sector",
        "P/E", "P/B", "Div Yield", "Mkt Cap (HKD)",
    ]
    write_header_row(ws, headers, row=4)

    widths = [12, 10, 12, 12, 14, 20, 20, 22, 10, 10, 12, 22]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[col(i)].width = w

    for r_idx, (_, row) in enumerate(holdings.iterrows(), start=5):
        ticker   = str(row["Ticker"]).strip()
        currency = str(row["Local Currency"]).strip()
        ticker, currency = TICKER_REMAP.get((ticker, currency), (ticker, currency))
        weight   = row["Weighting"]

        result = beta_map.get(ticker, {})
        n_days = result.get("n_days", "—")
        status = result.get("status", "Failed")
        if status == "OK" and "excel_link" in result:
            beta_val = result["excel_link"]
        else:
            beta_val = result.get("beta", "N/A")

        f = fundamentals_map.get(ticker, {})
        pe           = f.get("pe")
        pb           = f.get("pb")
        div_yield    = f.get("div_yield")
        mkt_hkd      = f.get("mkt_cap_hkd")
        gics_sector  = f.get("gics_sector")

        def _val(v):
            return v if v is not None else "N/A"

        w_num = pd.to_numeric(weight, errors="coerce")
        if pd.isna(w_num) or w_num == 0:
            long_short = "—"
        elif w_num > 0:
            long_short = "Long"
        else:
            long_short = "Short"

        vals  = [ticker, currency, weight, long_short, n_days, beta_val, status,
                 _val(gics_sector), _val(pe), _val(pb), _val(div_yield), _val(mkt_hkd)]
        fmts  = [None, None, "0%", None, "0", "0.0000", None,
                 None, "0.00", "0.00", '0.00%', "#,##0"]
        fonts = [DATA_FONT, DATA_FONT, DATA_FONT, DATA_FONT, DATA_FONT,
                 Font(name="Arial", size=9, bold=True), DATA_FONT,
                 DATA_FONT, DATA_FONT, DATA_FONT, DATA_FONT, DATA_FONT]

        for c_idx, (v, fmt, fnt) in enumerate(zip(vals, fmts, fonts), 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=v)
            cell.font      = fnt
            cell.border    = BORDER
            cell.alignment = Alignment(horizontal="center" if c_idx > 1 else "left")
            if fmt and isinstance(v, (int, float)):
                cell.number_format = fmt
            elif fmt and isinstance(v, str) and v.startswith("="):
                cell.number_format = fmt

        # Colour Long/Short cell
        ls_cell = ws.cell(row=r_idx, column=4)
        if long_short == "Long":
            ls_cell.fill = PatternFill("solid", start_color="FFE2EFDA")
            ls_cell.font = Font(name="Arial", size=9, color="FF375623")
        elif long_short == "Short":
            ls_cell.fill = PatternFill("solid", start_color="FFFFC7CE")
            ls_cell.font = Font(name="Arial", size=9, color="FF9C0006")
        else:
            ls_cell.font = DATA_FONT

        if status == "OK":
            ws.cell(row=r_idx, column=7).fill = PatternFill("solid", start_color="FFE2EFDA")
        elif status == "Failed":
            ws.cell(row=r_idx, column=7).fill = PatternFill("solid", start_color="FFFFC7CE")
        else:
            ws.cell(row=r_idx, column=7).fill = PatternFill("solid", start_color="FFFFEB9C")

    # ---------------------------------------------------------------------------
    # Weighted summary rows
    # ---------------------------------------------------------------------------
    n_holdings  = len(holdings)
    first_row   = 5
    last_row    = 4 + n_holdings
    summary_row = last_row + 2

    gold_fill = PatternFill("solid", start_color="FFFFF2CC")
    bold_font = Font(name="Arial", bold=True, size=9)

    # Col references
    w  = f"C{first_row}:C{last_row}"   # weights
    f_ = f"F{first_row}:F{last_row}"   # beta   (col F)
    pe = f"I{first_row}:I{last_row}"   # P/E    (col I)
    pb = f"J{first_row}:J{last_row}"   # P/B    (col J)

    # Weighted beta: SUMPRODUCT(weights, betas)
    weighted_beta_formula = f"=SUMPRODUCT({w},{f_})"

    # Weighted P/E (long only): numerator = Σ(w*pe where long & pe numeric)
    #                           denominator = Σ(w where long & pe numeric)
    weighted_pe_formula = (
        f"=IFERROR("
        f"SUMPRODUCT(({w}>0)*ISNUMBER({pe})*{w}*IFERROR({pe},0))"
        f"/SUMPRODUCT(({w}>0)*ISNUMBER({pe})*{w})"
        f',\"N/A\")'
    )

    # Weighted P/B (long only): same pattern with pb column
    weighted_pb_formula = (
        f"=IFERROR("
        f"SUMPRODUCT(({w}>0)*ISNUMBER({pb})*{w}*IFERROR({pb},0))"
        f"/SUMPRODUCT(({w}>0)*ISNUMBER({pb})*{w})"
        f',\"N/A\")'
    )

    summary_rows = [
        ("Weighted Beta (45-Day)",       weighted_beta_formula, "0.0000"),
        ("Weighted Avg P/E (Long only)", weighted_pe_formula,   "0.00"),
        ("Weighted Avg P/B (Long only)", weighted_pb_formula,   "0.00"),
    ]

    for i, (label, value, num_fmt) in enumerate(summary_rows):
        sr = summary_row + i

        lc = ws.cell(row=sr, column=1, value=label)
        lc.font      = bold_font
        lc.fill      = gold_fill
        lc.border    = BORDER
        lc.alignment = Alignment(horizontal="left")
        ws.merge_cells(start_row=sr, start_column=1, end_row=sr, end_column=4)

        display = value if value is not None else "N/A"
        vc = ws.cell(row=sr, column=5, value=display)
        vc.font          = bold_font
        vc.fill          = gold_fill
        vc.border        = BORDER
        vc.alignment     = Alignment(horizontal="center")
        vc.number_format = num_fmt
        ws.merge_cells(start_row=sr, start_column=5, end_row=sr, end_column=12)

    # ---------------------------------------------------------------------------
    # Long-side sector exposure
    # ---------------------------------------------------------------------------
    exposure      = _long_sector_exposure(holdings, fundamentals_map)
    exposure_start = summary_row + len(summary_rows) + 2

    blue_fill  = PatternFill("solid", start_color="FF1F4E79")
    light_fill = PatternFill("solid", start_color="FFDCE6F1")

    # Section header
    hdr = ws.cell(row=exposure_start, column=1, value="Long-Side Sector Exposure (% of Long Book)")
    hdr.font      = Font(name="Arial", bold=True, size=9, color="FFFFFFFF")
    hdr.fill      = blue_fill
    hdr.border    = BORDER
    hdr.alignment = Alignment(horizontal="left")
    ws.merge_cells(start_row=exposure_start, start_column=1,
                   end_row=exposure_start, end_column=12)

    # Sub-header row
    sub_row = exposure_start + 1
    for col_idx, sub_label in [(1, "GICS Sector"), (5, "# Tickers"), (7, "% of Long Book")]:
        sc = ws.cell(row=sub_row, column=col_idx, value=sub_label)
        sc.font      = Font(name="Arial", bold=True, size=9, color="FFFFFFFF")
        sc.fill      = blue_fill
        sc.border    = BORDER
        sc.alignment = Alignment(horizontal="center" if col_idx > 1 else "left")
    ws.merge_cells(start_row=sub_row, start_column=1, end_row=sub_row, end_column=4)
    ws.merge_cells(start_row=sub_row, start_column=5, end_row=sub_row, end_column=6)
    ws.merge_cells(start_row=sub_row, start_column=7, end_row=sub_row, end_column=12)

    for j, (sector, n_tickers, pct) in enumerate(exposure, start=1):
        sr = exposure_start + 1 + j
        row_fill = light_fill if j % 2 == 0 else PatternFill()

        sc = ws.cell(row=sr, column=1, value=sector)
        sc.font      = DATA_FONT
        sc.fill      = row_fill
        sc.border    = BORDER
        sc.alignment = Alignment(horizontal="left")
        ws.merge_cells(start_row=sr, start_column=1, end_row=sr, end_column=4)

        nc = ws.cell(row=sr, column=5, value=n_tickers)
        nc.font          = DATA_FONT
        nc.fill          = row_fill
        nc.border        = BORDER
        nc.number_format = "0"
        nc.alignment     = Alignment(horizontal="center")
        ws.merge_cells(start_row=sr, start_column=5, end_row=sr, end_column=6)

        pc = ws.cell(row=sr, column=7, value=pct)
        pc.font          = Font(name="Arial", bold=True, size=9)
        pc.fill          = row_fill
        pc.border        = BORDER
        pc.number_format = "0%"
        pc.alignment     = Alignment(horizontal="center")
        ws.merge_cells(start_row=sr, start_column=7, end_row=sr, end_column=12)

    ws.freeze_panes = "A5"

# ---------------------------------------------------------------------------
# MAIN
# ---------------------------------------------------------------------------

def build_audit_workbook(holdings: pd.DataFrame) -> Workbook:
    """Build the Beta Calculation Audit workbook from a holdings DataFrame.

    Pure in-memory: takes the parsed holdings and returns an openpyxl Workbook.
    Callers (the CLI runner or the web app) handle reading input / saving output.
    """
    holdings = holdings.copy()
    holdings["Weighting"] = pd.to_numeric(holdings["Weighting"], errors="coerce")

    print("\nFetching Benchmark: 2800.HK (Hang Seng ETF)...")
    benchmark_raw = get_prices("2800", "HKD")
    if benchmark_raw is None or benchmark_raw.empty:
        raise RuntimeError("Could not fetch benchmark (2800.HK). Cannot continue.")

    benchmark = benchmark_raw.tail(60)

    wb = Workbook()
    del wb[wb.sheetnames[0]]

    create_fx_sheet(wb)

    beta_map         = {}
    fundamentals_map = {}

    for _, row in holdings.iterrows():
        ticker   = str(row["Ticker"]).strip()
        currency = str(row["Local Currency"]).strip()

        print(f"\nProcessing: {ticker} ({currency})...")

        # --- Prices & Beta ---
        asset = get_prices(ticker, currency)

        if asset is None or asset.empty:
            print(f"  -> No price data for {ticker}, skipping beta.")
            resolved_t, _ = TICKER_REMAP.get((ticker, currency), (ticker, currency))
            beta_map[resolved_t] = {"beta": "N/A", "status": "Failed", "n_days": 0}
        else:
            combined = align_to_benchmark(asset, benchmark)

            if len(combined) < 46:
                msg = f"Insufficient aligned data ({len(combined)} days)"
                print(f"  -> {msg}")
                beta_map[ticker] = {"beta": msg, "status": "Insufficient data",
                                    "n_days": len(combined)}
            else:
                returns  = combined.pct_change().dropna()
                last     = returns.tail(45)
                cov      = last[["asset_close", "bm_close"]].cov().iloc[0, 1]
                var      = last["bm_close"].var()
                beta_val = cov / var if var != 0 else float("nan")

                t, c_ = TICKER_REMAP.get((ticker, currency), (ticker, currency))
                sheet_name  = t[:31].replace("/", "_").replace("\\", "_")
                beta_result = {
                    "covariance": round(cov, 8),
                    "variance":   round(var, 8),
                    "beta":       round(beta_val, 6),
                    "n_days":     len(combined),
                    "status":     "OK",
                    "excel_link": f"='{sheet_name}'!B{len(combined) + 6}",
                }
                beta_map[t] = beta_result
                write_stock_sheet(wb, ticker, currency, combined, beta_result)

        # --- Fundamentals ---
        print(f"  Fetching fundamentals for {ticker}...")
        fund = get_fundamentals(ticker, currency)
        resolved_ticker, _ = TICKER_REMAP.get((ticker, currency), (ticker, currency))
        fundamentals_map[resolved_ticker] = fund
        print(f"  -> PE={fund['pe']}, PB={fund['pb']}, "
              f"DivYield={fund['div_yield']}, "
              f"MktCap(HKD)={fund['mkt_cap_hkd']}")

    write_summary_sheet(wb, holdings, beta_map, fundamentals_map)
    write_fundamentals_sheet(wb, holdings, fundamentals_map)

    sheets = wb.sheetnames
    wb.move_sheet("Summary",      offset=-(len(sheets) - 2))
    wb.move_sheet("Fundamentals", offset=-(len(sheets) - 2))

    return wb


def main(input_path: str = input_file,
         output_path: str = "Beta_Calculation_Audit.xlsx") -> str:
    """CLI entry point: read holdings from disk, build the workbook, save it."""
    holdings = pd.read_excel(input_path)
    wb = build_audit_workbook(holdings)
    wb.save(output_path)
    print(f"\nDone. Audit file saved: {output_path}")
    return output_path


if __name__ == "__main__":
    main()
